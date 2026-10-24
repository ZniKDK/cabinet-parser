"""Excel на входе (список артикулов) и на выходе (результат с продолжением после сбоя)."""

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Эти статусы повторно не проверяем: товар уже собран или его точно нет на сайте
FINAL_STATUSES = ("ok", "не найден")


class ExcelError(Exception):
    """Проблема с входным или выходным файлом; текст понятен человеку."""


def _norm(text) -> str:
    return str(text or "").strip().casefold()


def _sku_str(value) -> str:
    """Excel хранит числа как float: 777.0 → «777»."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_skus(path: Path, column: str, sheet: str | None = None) -> list[str]:
    """Читает артикулы из колонки с заголовком column. Пустые и повторы убирает."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError as error:
        raise ExcelError(f"Файл не найден: {path}") from error
    ws = wb[sheet] if sheet else wb.active

    rows = ws.iter_rows(values_only=True)
    header = next(rows, ())
    try:
        index = [_norm(cell) for cell in header].index(_norm(column))
    except ValueError:
        raise ExcelError(f"В первой строке нет колонки «{column}». Есть: "
                         + ", ".join(str(c) for c in header if c)) from None

    skus, seen = [], set()
    for row in rows:
        value = row[index] if index < len(row) else None
        if value is None or str(value).strip() == "":
            continue
        sku = _sku_str(value)
        if sku not in seen:
            seen.add(sku)
            skus.append(sku)
    wb.close()
    return skus


class ResultWriter:
    """Пишет результат в Excel. Если файл уже есть — дописывает и знает, что готово."""

    def __init__(self, path: Path, columns: list[tuple[str, str]]) -> None:
        self.path = Path(path)
        self.columns = columns            # (имя поля, заголовок колонки)
        self.rows: dict[str, list] = {}   # артикул → строка; порядок вставки сохраняется
        self.done: set[str] = set()
        if self.path.exists():
            self._load()

    def _headers(self) -> list[str]:
        return ["Артикул", *[title for _, title in self.columns], "Статус", "Ссылка", "Проверено"]

    def _load(self) -> None:
        ws = load_workbook(self.path).active
        status_index = len(self.columns) + 1
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            sku = str(row[0])
            self.rows[sku] = list(row)
            if row[status_index] in FINAL_STATUSES:
                self.done.add(sku)

    def add(self, sku: str, values: dict, status: str, url: str) -> None:
        row = [sku]
        for name, _ in self.columns:
            value = values.get(name)
            # Decimal в Excel пишем числом, чтобы с ним можно было считать
            row.append(float(value) if isinstance(value, Decimal) else value)
        row += [status, url, datetime.now().strftime("%d.%m.%Y %H:%M")]
        self.rows[sku] = row
        if status in FINAL_STATUSES:
            self.done.add(sku)

    def save(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Результат"
        ws.append(self._headers())
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in self.rows.values():
            ws.append(row)
        # Числа — с разделителями тысяч и двумя знаками
        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        for index, header in enumerate(self._headers(), 1):
            ws.column_dimensions[get_column_letter(index)].width = max(12, len(header) + 4)
        ws.freeze_panes = "A2"
        self.path.parent.mkdir(parents=True, exist_ok=True)
        # Сначала во временный файл: если упадём посреди записи, старый результат цел
        tmp = self.path.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        tmp.replace(self.path)
