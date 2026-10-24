"""Тесты чтения артикулов из Excel и записи результата с продолжением."""

from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from cabinet_parser.excel_io import ExcelError, ResultWriter, read_skus


def make_input(tmp_path, rows, header="Артикул"):
    wb = Workbook()
    ws = wb.active
    ws.append(["Название", header])
    for row in rows:
        ws.append(["товар", row])
    path = tmp_path / "in.xlsx"
    wb.save(path)
    return path


def test_read_skus_cleans_and_dedupes(tmp_path):
    path = make_input(tmp_path, ["BG1001", " BG1002 ", None, "BG1001", 12345, 777.0])
    assert read_skus(path, "Артикул") == ["BG1001", "BG1002", "12345", "777"]


def test_column_not_found(tmp_path):
    path = make_input(tmp_path, ["BG1001"], header="Код")
    with pytest.raises(ExcelError, match="Артикул"):
        read_skus(path, "Артикул")


def test_column_match_ignores_case_and_spaces(tmp_path):
    path = make_input(tmp_path, ["BG1001"], header=" артикул ")
    assert read_skus(path, "Артикул") == ["BG1001"]


def test_writer_writes_and_resumes(tmp_path):
    out = tmp_path / "out.xlsx"
    columns = [("price", "Цена"), ("stock", "Наличие")]
    writer = ResultWriter(out, columns)
    writer.add("BG1001", {"price": Decimal("6420.5"), "stock": "В магазине"}, "ok", "https://x/1")
    writer.add("BG9999", {}, "не найден", "")
    writer.add("BG1002", {}, "ошибка: таймаут", "")
    writer.save()

    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == ["Артикул", "Цена", "Наличие", "Статус", "Ссылка", "Проверено"]
    assert ws.cell(2, 2).value == 6420.5
    assert ws.cell(2, 2).number_format == "#,##0.00"

    # При повторном запуске готовые и ненайденные пропускаются, ошибки повторяются
    again = ResultWriter(out, columns)
    assert again.done == {"BG1001", "BG9999"}
    again.add("BG1002", {"price": Decimal("4380")}, "ok", "https://x/2")
    again.save()
    rows = list(load_workbook(out).active.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["BG1001", "BG9999", "BG1002"]
    assert rows[2][3] == "ok"
