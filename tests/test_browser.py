"""Сквозные тесты с настоящим браузером (Chromium) на демо-магазине."""

import socket
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cabinet_parser.browser import LoginError
from cabinet_parser.config import load_config
from cabinet_parser.runner import run
from demo_shop.server import DEMO_LOGIN, DEMO_PASSWORD, start

pytestmark = pytest.mark.browser
EXAMPLE = Path(__file__).parent.parent / "examples" / "demo_shop.yaml"


def free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


@pytest.fixture(scope="module")
def site(tmp_path_factory):
    """Демо-магазин на свободном порту и описание сайта под этот порт."""
    port = free_port()
    server = start(port)
    text = EXAMPLE.read_text(encoding="utf-8").replace("8765", str(port))
    text = text.replace("delay_sec: 1 ", "delay_sec: 0.5 ")
    path = tmp_path_factory.mktemp("site") / "demo.yaml"
    path.write_text(text, encoding="utf-8")
    yield load_config(path)
    server.shutdown()


def make_input(tmp_path, skus):
    wb = Workbook()
    ws = wb.active
    ws.append(["Артикул"])
    for sku in skus:
        ws.append([sku])
    path = tmp_path / "in.xlsx"
    wb.save(path)
    return path


def rows(path):
    return {r[0]: r for r in load_workbook(path).active.iter_rows(min_row=2, values_only=True)}


def test_full_run_and_resume(site, tmp_path):
    source = make_input(tmp_path, ["BG1001", "BG1004", "BG9999", "BG1002"])
    out = tmp_path / "out.xlsx"
    state = tmp_path / "session.json"
    report = run(site, source, out, state, DEMO_LOGIN, DEMO_PASSWORD)
    assert (report.ok, report.not_found, report.errors) == (3, 1, 0)

    result = rows(out)
    # Колонки: Артикул, Название, Цена, Наличие, Количество, Срок, Статус, Ссылка, Проверено
    assert result["BG1001"][1] == "Дрель-шуруповёрт Berger 18 В"
    assert result["BG1001"][2] == 6420.5          # цена из кабинета, не розничная 7990
    assert result["BG1001"][3] == "В магазине"
    assert result["BG1004"][4] == 0
    assert result["BG1002"][5] == "3–5 дней"
    assert result["BG9999"][6] == "не найден"
    assert result["BG1001"][7].endswith("/product/BG1001")
    assert state.exists()

    # Повторный запуск: всё готово, браузер даже ничего не ищет
    again = run(site, source, out, state, DEMO_LOGIN, DEMO_PASSWORD)
    assert (again.skipped, again.ok) == (4, 0)


def test_input_columns_copied_to_result(site, tmp_path):
    wb = Workbook()
    wb.active.append(["Название в прайсе", "Артикул"])
    wb.active.append(["Дрель из прайса Berger", "BG1001"])
    source = tmp_path / "prices.xlsx"
    wb.save(source)
    cfg = replace(site, input_keep=["Название в прайсе"])
    run(cfg, source, tmp_path / "out.xlsx", tmp_path / "s.json", DEMO_LOGIN, DEMO_PASSWORD)
    row = rows(tmp_path / "out.xlsx")["BG1001"]
    assert row[1] == "Дрель из прайса Berger"       # колонка из входного файла
    assert row[2] == "Дрель-шуруповёрт Berger 18 В"  # название с сайта


def test_saved_session_works_without_password(site, tmp_path):
    state = tmp_path / "session.json"
    run(site, make_input(tmp_path, ["BG1005"]), tmp_path / "a.xlsx", state, DEMO_LOGIN, DEMO_PASSWORD)
    # Пароля нет, но сохранённая сессия действует
    report = run(site, make_input(tmp_path, ["BG1003"]), tmp_path / "b.xlsx", state, "", "")
    assert report.ok == 1
    assert rows(tmp_path / "b.xlsx")["BG1003"][2] == 9870


def test_wrong_password(site, tmp_path):
    with pytest.raises(LoginError, match="Вход не удался"):
        run(site, make_input(tmp_path, ["BG1001"]), tmp_path / "out.xlsx",
            tmp_path / "s.json", DEMO_LOGIN, "wrong")


def test_no_credentials_and_no_session(site, tmp_path):
    with pytest.raises(LoginError, match="CABINET_LOGIN"):
        run(site, make_input(tmp_path, ["BG1001"]), tmp_path / "out.xlsx", tmp_path / "s.json", "", "")
